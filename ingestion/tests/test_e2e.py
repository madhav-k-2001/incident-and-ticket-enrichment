import os
import pytest
from reportlab.pdfgen import canvas
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.main import app
from app.db.database import SessionLocal
from app.db.models import Document, DocumentChunk
from worker.worker import IngestionWorker


@pytest.fixture
def test_files(tmp_path):
    # 1. Create a PDF about Gemini Embedding
    pdf_path = tmp_path / "gemini_guide.pdf"
    c = canvas.Canvas(str(pdf_path))
    c.drawString(100, 750, "Gemini Embedding 2 is Google's multimodal embedding model.")
    c.drawString(100, 720, "It supports Matryoshka Representation Learning with 768 dimensions.")
    c.showPage()
    c.drawString(100, 750, "Rate limits must be respected: 100 RPM, 30K TPM, and 1K RPD.")
    c.showPage()
    c.save()

    # 2. Create a DOCX about PostgreSQL pgvector
    docx_path = tmp_path / "pgvector_notes.docx"
    doc = DocxDocument()
    doc.add_heading("PostgreSQL and pgvector", level=1)
    doc.add_paragraph("pgvector enables vector similarity search directly in Postgres.")
    doc.add_paragraph("It supports HNSW and IVFFlat indexes with cosine distance operations.")
    doc.save(str(docx_path))

    return str(pdf_path), str(docx_path)


def test_end_to_end_pipeline(test_files):
    pdf_file, docx_file = test_files
    client = TestClient(app)

    # 1. Upload both files
    with open(pdf_file, "rb") as f_pdf, open(docx_file, "rb") as f_docx:
        upload_resp = client.post(
            "/api/upload",
            files=[
                ("files", ("gemini_guide.pdf", f_pdf.read(), "application/pdf")),
                (
                    "files",
                    (
                        "pgvector_notes.docx",
                        f_docx.read(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                ),
            ],
        )

    assert upload_resp.status_code == 200
    upload_data = upload_resp.json()
    assert len(upload_data["documents"]) == 2

    pdf_doc_id = upload_data["documents"][0]["id"]
    docx_doc_id = upload_data["documents"][1]["id"]

    # 2. Process documents using IngestionWorker
    worker = IngestionWorker()
    worker.process_document(pdf_doc_id)
    worker.process_document(docx_doc_id)

    # 3. Verify in database
    db = SessionLocal()
    try:
        pdf_doc = db.query(Document).filter(Document.id == pdf_doc_id).first()
        assert pdf_doc is not None
        assert pdf_doc.status == "COMPLETED"
        assert pdf_doc.total_chunks > 0
        assert pdf_doc.processed_chunks == pdf_doc.total_chunks

        docx_doc = db.query(Document).filter(Document.id == docx_doc_id).first()
        assert docx_doc is not None
        assert docx_doc.status == "COMPLETED"
        assert docx_doc.total_chunks > 0

        # Check chunks in pgvector
        pdf_chunks = db.query(DocumentChunk).filter(DocumentChunk.document_id == pdf_doc_id).all()
        assert len(pdf_chunks) == pdf_doc.total_chunks
        for chunk in pdf_chunks:
            assert len(chunk.embedding) == 768

        docx_chunks = db.query(DocumentChunk).filter(DocumentChunk.document_id == docx_doc_id).all()
        assert len(docx_chunks) == docx_doc.total_chunks
    finally:
        db.close()

    # 4. Verify Semantic Vector Search
    search_resp = client.post("/api/search", json={"query": "Gemini Embedding 2 dimensions", "top_k": 3})
    assert search_resp.status_code == 200
    search_data = search_resp.json()
    assert search_data["results_count"] > 0
    top_result = search_data["results"][0]
    assert "similarity" in top_result
    assert "content" in top_result

    # 5. Verify Quota Meter Endpoint
    quota_resp = client.get("/api/quota")
    assert quota_resp.status_code == 200
    quota_data = quota_resp.json()
    assert "rpm" in quota_data
    assert "tpm" in quota_data
    assert "rpd" in quota_data

    # Clean up test documents
    client.delete(f"/api/documents/{pdf_doc_id}")
    client.delete(f"/api/documents/{docx_doc_id}")
