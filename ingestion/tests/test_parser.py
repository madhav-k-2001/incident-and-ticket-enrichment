import os
import pytest
from reportlab.pdfgen import canvas
from docx import Document as DocxDocument

from worker.parser import parse_pdf, parse_docx, parse_document


@pytest.fixture
def sample_pdf(tmp_path):
    pdf_path = tmp_path / "sample.pdf"
    c = canvas.Canvas(str(pdf_path))
    # Page 1
    c.drawString(100, 750, "Hello World from PDF Page 1.")
    c.drawString(100, 700, "This is second line of text.")
    c.showPage()
    # Page 2
    c.drawString(100, 750, "Content on PDF Page 2.")
    c.showPage()
    c.save()
    return str(pdf_path)


@pytest.fixture
def sample_docx(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_heading("Document Title", level=1)
    doc.add_paragraph("This is the first paragraph of test document.")
    doc.add_paragraph("Here is another paragraph containing relevant data.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Key1"
    table.cell(0, 1).text = "Val1"
    table.cell(1, 0).text = "Key2"
    table.cell(1, 1).text = "Val2"
    doc.save(str(docx_path))
    return str(docx_path)


def test_parse_pdf(sample_pdf):
    results = parse_pdf(sample_pdf)
    assert len(results) == 2
    assert results[0]["page_number"] == 1
    assert "Page 1" in results[0]["text"]
    assert results[1]["page_number"] == 2
    assert "Page 2" in results[1]["text"]


def test_parse_docx(sample_docx):
    results = parse_docx(sample_docx)
    assert len(results) >= 1
    full_text = " ".join(r["text"] for r in results)
    assert "Document Title" in full_text
    assert "first paragraph" in full_text
    assert "Key1 | Val1" in full_text


def test_parse_document_dispatcher(sample_pdf, sample_docx, tmp_path):
    res_pdf = parse_document(sample_pdf)
    assert len(res_pdf) == 2

    res_docx = parse_document(sample_docx)
    assert len(res_docx) >= 1

    unsupported = tmp_path / "test.txt"
    unsupported.write_text("plain text")
    with pytest.raises(ValueError) as excinfo:
        parse_document(str(unsupported))
    assert "Unsupported file type" in str(excinfo.value)
