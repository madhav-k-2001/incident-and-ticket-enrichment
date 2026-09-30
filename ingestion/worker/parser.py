import os
import re
from pathlib import Path
from typing import List, Dict, Any
import logging
from pypdf import PdfReader
from docx import Document as DocxDocument

logger = logging.getLogger(__name__)


def parse_pdf(file_path: str) -> List[Dict[str, Any]]:
    """
    Extracts text from a PDF file page by page.
    Returns:
        List of dicts: [{"page_number": int, "text": str}]
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"PDF file not found at {file_path}")

    results = []
    reader = PdfReader(file_path)

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("PDF is encrypted and password-protected.")

    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            results.append({
                "page_number": idx + 1,
                "text": text
            })

    if not results:
        logger.warning(f"No extractable text found in PDF: {file_path}")

    return results


def parse_docx(file_path: str) -> List[Dict[str, Any]]:
    """
    Extracts text from a DOCX file, parsing paragraphs and tables.
    Returns:
        List of dicts: [{"page_number": int, "text": str}]
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"DOCX file not found at {file_path}")

    doc = DocxDocument(file_path)
    text_blocks = []

    # Extract paragraphs
    for p in doc.paragraphs:
        p_text = p.text.strip()
        if p_text:
            text_blocks.append(p_text)

    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_blocks.append(row_text)

    if not text_blocks:
        logger.warning(f"No extractable text found in DOCX: {file_path}")
        return []

    # DOCX does not define static physical pages, so divide into pseudo-pages
    # of ~3000 characters (approx 1 page of text) for logical grouping
    results = []
    current_page_text = []
    current_char_count = 0
    page_number = 1

    for block in text_blocks:
        current_page_text.append(block)
        current_char_count += len(block)

        if current_char_count >= 3000:
            results.append({
                "page_number": page_number,
                "text": "\n\n".join(current_page_text)
            })
            page_number += 1
            current_page_text = []
            current_char_count = 0

    if current_page_text:
        results.append({
            "page_number": page_number,
            "text": "\n\n".join(current_page_text)
        })

    return results


_HEADING_RE = re.compile(r"^(#{1,6})[ \t]+(.+?)[ \t]*#*[ \t]*$")
_FENCE_RE = re.compile(r"^\s{0,3}(```|~~~)")


def parse_markdown(file_path: str) -> List[Dict[str, Any]]:
    """
    Parses a Markdown file into sections split on ATX headings (# .. ######).
    Headings inside fenced code blocks are ignored. Text before the first
    heading becomes its own section.
    Returns:
        List of dicts: [{"page_number": int (1-based section number),
                         "text": str, "heading": str (e.g. "Intro > Setup")}]
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Markdown file not found at {file_path}")

    with open(file_path, "r", encoding="utf-8-sig", errors="replace") as f:
        lines = f.read().splitlines()

    sections: List[Dict[str, Any]] = []
    stack: List[tuple] = []  # (level, title)
    current_heading = ""
    current_lines: List[str] = []
    in_fence = False

    def flush():
        text = "\n".join(current_lines).strip()
        # A heading with no body still carries no content worth embedding
        has_body = any(l.strip() and not _HEADING_RE.match(l) for l in current_lines)
        if text and has_body:
            sections.append({
                "page_number": len(sections) + 1,
                "text": text,
                "heading": current_heading,
            })

    for line in lines:
        if _FENCE_RE.match(line):
            in_fence = not in_fence
        m = None if in_fence else _HEADING_RE.match(line)
        if m:
            flush()
            level, title = len(m.group(1)), m.group(2).strip()
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, title))
            current_heading = " > ".join(t for _, t in stack)
            current_lines = [line]
        else:
            current_lines.append(line)
    flush()

    if not sections:
        logger.warning(f"No extractable text found in Markdown: {file_path}")

    return sections


def parse_document(file_path: str) -> List[Dict[str, Any]]:
    """
    Dispatcher based on file extension.
    Supports .pdf, .docx and .md.
    """
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return parse_docx(file_path)
    elif ext in [".md", ".markdown"]:
        return parse_markdown(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Only PDF, DOCX and Markdown files are supported.")
